import sqlite3
from contextlib import closing
from io import BytesIO
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook
from sqlalchemy import create_engine, func, inspect, select

from app.db.migrations import migrate_history_ownership
from app.db.session import get_db
from app.importers.excel import COLUMNS
from app.main import app
from app.models.entities import InferenceLog


def user_headers():
    return {"X-Demo-Role": "user", "X-Demo-User": str(uuid4())}


def save_session(client, headers=None):
    response = client.post(
        "/api/consultations", json={"facts": {"purpose": "office"}}, headers=headers
    )
    assert response.status_code == 201, response.text
    return response.json()["session_id"]


def test_history_scoped_by_account_including_detail_pagination_dashboard(client):
    alice, bob = user_headers(), user_headers()
    admin_id = save_session(client)
    alice_ids = [save_session(client, alice) for _ in range(2)]
    bob_id = save_session(client, bob)
    for headers, expected in [({}, [admin_id]), (alice, alice_ids[::-1]), (bob, [bob_id])]:
        history = client.get("/api/consultations", headers=headers).json()
        assert [row["id"] for row in history["items"]] == expected
        assert history["total"] == len(expected)
        dashboard = client.get("/api/dashboard", headers=headers).json()
        assert dashboard["consultation_sessions"] == len(expected)
        assert [s["id"] for s in dashboard["recent_consultations"]] == expected
        for session_id in [admin_id, *alice_ids, bob_id]:
            response = client.get(f"/api/consultations/{session_id}", headers=headers)
            assert response.status_code == (200 if session_id in expected else 404)
    second_page = client.get("/api/consultations?page=2&page_size=1", headers=alice).json()
    assert second_page["total"] == 2
    assert [s["id"] for s in second_page["items"]] == [alice_ids[0]]


def test_history_requires_explicit_valid_identity(client):
    del client.headers["X-Demo-Role"]
    for path in ["/api/consultations", "/api/consultations/1", "/api/dashboard"]:
        assert client.get(path).status_code == 422
    assert client.post("/api/consultations", json={"facts": {}}).status_code == 422
    for headers in [
        {"X-Demo-Role": "user"},
        {"X-Demo-Role": "user", "X-Demo-User": "invalid"},
        {"X-Demo-Role": "invalid"},
    ]:
        assert client.get("/api/consultations", headers=headers).status_code == 422


def test_admin_bulk_history_delete_is_atomic_owned_and_cascades(client):
    first, second = save_session(client), save_session(client)
    user = user_headers()
    other = save_session(client, user)
    for invalid, status in [([], 422), ([0], 422), ([first, 999999], 404), ([first, other], 404)]:
        response = client.request("DELETE", "/api/consultations", json={"ids": invalid})
        assert response.status_code == status
        assert client.get(f"/api/consultations/{first}").status_code == 200
    assert (
        client.request(
            "DELETE", "/api/consultations", json={"ids": [other]}, headers=user
        ).status_code
        == 403
    )
    db_generator = app.dependency_overrides[get_db]()
    with closing(db_generator):
        db = next(db_generator)
        assert (
            db.scalar(
                select(func.count())
                .select_from(InferenceLog)
                .where(InferenceLog.session_id.in_([first, second]))
            )
            > 0
        )
        response = client.request(
            "DELETE", "/api/consultations", json={"ids": [first, second, first]}
        )
        assert response.status_code == 200
        assert response.json() == {"deleted": 2}
        assert (
            db.scalar(
                select(func.count())
                .select_from(InferenceLog)
                .where(InferenceLog.session_id.in_([first, second]))
            )
            == 0
        )
        assert (
            db.scalar(
                select(func.count())
                .select_from(InferenceLog)
                .where(InferenceLog.session_id == other)
            )
            > 0
        )
    assert client.get("/api/consultations").json()["total"] == 0
    assert client.get(f"/api/consultations/{other}", headers=user).status_code == 200
    assert client.get("/api/dashboard").json()["total_rules"] == 30


def test_excel_template_headers_only_match_original_and_can_be_filled(client):
    response = client.get("/api/products/template")
    assert response.status_code == 200
    assert (
        response.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert 'filename="Laptop_template.xlsx"' in response.headers["content-disposition"]
    workbook = load_workbook(BytesIO(response.content))
    assert len(workbook.worksheets) == 1
    sheet = workbook.active
    assert sheet.max_row == 1
    assert sheet.max_column == 12
    assert [cell.value for cell in sheet[1]] == COLUMNS
    source = load_workbook(Path(__file__).resolve().parents[2] / "Laptop_data.xlsx", read_only=True)
    assert list(next(source.active.values)) == COLUMNS
    source.close()
    sheet.append(
        [
            1,
            "TEMPLATE-001",
            "Template laptop",
            20000000,
            "Brand",
            "CPU",
            "GPU",
            "16 GB",
            "1.5 kg",
            "512 GB SSD",
            "Silver",
            "15.6 inch 144 Hz",
        ]
    )
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    report = client.post("/api/products/import", files={"file": ("filled.xlsx", output.getvalue())})
    assert report.status_code == 200
    assert report.json()["created"] == 1 and report.json()["failed"] == 0


def test_existing_database_upgrade_keeps_snapshots_and_backs_up_once(tmp_path):
    path = tmp_path / "legacy.db"
    with closing(sqlite3.connect(path)) as db:
        db.executescript("""
            CREATE TABLE consultation_sessions (
                id INTEGER PRIMARY KEY, created_at TEXT, initial_facts_json JSON,
                final_facts_json JSON, recommendation_json JSON
            );
            INSERT INTO consultation_sessions VALUES (8, '2026-01-01', '{}', '{}', '{"steps": [1, 2]}');
            CREATE TABLE inference_logs (id INTEGER PRIMARY KEY, session_id INTEGER REFERENCES consultation_sessions(id));
            INSERT INTO inference_logs VALUES (12, 8);
            CREATE TABLE products (id INTEGER PRIMARY KEY, product_code TEXT);
            INSERT INTO products VALUES (51, 'PRESERVED');
        """)
    engine = create_engine(f"sqlite:///{path.as_posix()}")
    try:
        migrate_history_ownership(engine)
        migrate_history_ownership(engine)
        with closing(sqlite3.connect(path)) as db:
            assert db.execute(
                "SELECT id, owner_id, recommendation_json FROM consultation_sessions"
            ).fetchall() == [(8, "demo-admin", '{"steps": [1, 2]}')]
            assert db.execute("SELECT * FROM inference_logs").fetchall() == [(12, 8)]
            assert db.execute("SELECT * FROM products").fetchall() == [(51, "PRESERVED")]
        backups = list(tmp_path.glob("legacy.before-history-ownership-*.db"))
        assert len(backups) == 1
        with closing(sqlite3.connect(backups[0])) as backup:
            assert "owner_id" not in [
                r[1] for r in backup.execute("PRAGMA table_info(consultation_sessions)")
            ]
            assert backup.execute("SELECT id FROM consultation_sessions").fetchall() == [(8,)]
        assert "ix_consultation_sessions_owner_id" in {
            i["name"] for i in inspect(engine).get_indexes("consultation_sessions")
        }
    finally:
        engine.dispose()
